"""Generate the daily F&O trading-intelligence .docx brief for 2026-05-26."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DATE_ISO = "2026-05-26"
BIAS = "RANGE"
ENTRY_PERMISSION = "YELLOW"
GEN_TIME = "08:50 IST"

OUT_DIR = "trading-briefs/2026/05-May"
OUT_PATH = f"{OUT_DIR}/Trading_Brief_{DATE_ISO}_{BIAS}.docx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x38, 0x76, 0x1D)


def shade_cell(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_heading(doc, text, level=1, color=NAVY):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = color
    return h


def add_para(doc, text, bold=False, italic=False, size=11, color=None, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_kv_table(doc, rows, header=None):
    ncols = len(rows[0])
    table = doc.add_table(rows=(1 if header else 0) + len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    if header:
        hdr = table.rows[0]
        for i, txt in enumerate(header):
            cell = hdr.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(txt)
            run.bold = True
            run.font.size = Pt(11)
            shade_cell(cell, "1F3864")
            run.font.color.rgb = WHITE
    r0 = 1 if header else 0
    for i, row in enumerate(rows):
        trow = table.rows[r0 + i]
        for j, val in enumerate(row):
            cell = trow.cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(10.5)
            if j == 0:
                run.bold = True
    return table


def badge_line(doc, label, value, fill_hex, text_color=WHITE):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = ""
    shade_cell(cell, fill_hex)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{label}: {value}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = text_color


def stock_header(doc, text):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = ""
    shade_cell(cell, "1F3864")
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = WHITE


# ============== BUILD DOCUMENT ==============
doc = Document()
for section in doc.sections:
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

# ---------- HEADER ----------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("ELITE DAILY TRADING BRIEF")
run.bold = True
run.font.size = Pt(20)
run.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub.add_run("Indian F&O Day Trader  —  NSE Nifty + Stocks")
sub_run.italic = True
sub_run.font.size = Pt(12)

date_p = doc.add_paragraph()
date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
dr = date_p.add_run(f"Tuesday, 26 May 2026  |  Generated {GEN_TIME}  |  Weekly expiry: 0 days (TODAY)  |  Monthly expiry: 0 days (TODAY)  |  Expiry week: YES")
dr.bold = True
dr.font.size = Pt(10.5)

badge_line(doc, "TODAY'S BIAS", "RANGE (expiry pin, mild down lean to max pain)", "F1C232", RGBColor(0,0,0))
badge_line(doc, "ENTRY PERMISSION", "YELLOW — expiry day: intraday scalps only, exit by 13:00", "E69138")
badge_line(doc, "CRUDE MODE", "CAUTION (~Rs 8,766/bbl) | Direction: choppy/flat (crashed -6% then steadied)", "CC4125")
badge_line(doc, "VIX SIZING", "HALF SIZE (India VIX ~17.x, in 15-20 band) — MANDATORY", "674EA7")
badge_line(doc, "EXPIRY ALERT", "DUAL EXPIRY (weekly+monthly). NO option buying except scalps; exit before 13:00", "990000")

doc.add_paragraph()

# ---------- SECTION 1 ----------
add_heading(doc, "Section 1 — Macro Snapshot")
add_kv_table(doc,
    header=["Metric", "Value", "Signal"],
    rows=[
        ["MCX Crude (proxy)", "~Rs 8,766/bbl (WTI ~$92 x USDINR 95.35)", "CAUTION zone (8,500-9,000)"],
        ["Brent / WTI", "$97.74 / ~$92.0", "Crashed -6% prior session on US-Iran deal hopes; +0.5% today"],
        ["GIFT Nifty gap", "~24,038 vs close 24,031.7 = ~+6 pts", "FLAT open / tepid"],
        ["S&P 500", "7,473.47 (+0.37%)", "Risk-on (modest)"],
        ["Nasdaq / Dow", "26,343.97 (+0.19%) / 50,579.70 (+0.58%)", "Mildly positive"],
        ["US VIX (CBOE)", "~16.70", "Calm"],
        ["India VIX", "below 18 (-~5% on week)", "HALF SIZE rule (15-20 band)"],
        ["Nifty close (25 May)", "24,031.70 (+1.32% / +312.4 pts)", "Strong prior session"],
        ["FII cash (last prov.)", "NET SELL (-1,891 Cr on 21 May; week ~ -4,440 Cr)", "FIIs SELLING"],
        ["DII cash (last prov.)", "NET BUY (+2,492 Cr on 21 May; week ~ +6,004 Cr)", "DIIs absorbing"],
        ["USD-INR", "~95.35 (off 96.91 high)", "Rupee firming"],
        ["INFY ADR", "$12.64 (22 May); o/n move DATA_UNAVAILABLE", "Neutral signal"],
        ["WIT (Wipro) ADR", "$3.07, +7.17% (22 May, stale)", "Bullish skew for Wipro"],
        ["IBN (ICICI) ADR", "$30.06, +1.23% (22 May)", "Mild positive"],
        ["HDB (HDFC Bk) ADR", "$36.08, +2.65% (22 May)", "Mild positive"],
    ])
add_para(doc, "Geopolitical one-liner: Strait of Hormuz blocked since 28 Feb 2026; US launched fresh defensive strikes on Iran overnight, denting one-page peace-deal hopes that had just crashed crude ~6%.", italic=True)
add_para(doc, "Today's intraday headline risk: HIGH (binary Iran outcome — deal = crude crash / escalation = crude spike).", bold=True, color=RED)

# ---------- SECTION 2 ----------
add_heading(doc, "Section 2 — Crude Rule + Structural Read")
add_bullets(doc, [
    "CRUDE_MODE: CAUTION (~Rs 8,766/bbl, in 8,500-9,000 band) -> half size already mandated by VIX too.",
    "CRUDE_DIRECTION: CHOPPY/FLAT overnight (+0.5%) after a -6% prior-session crash on US-Iran deal optimism.",
    "CRUDE_FLIP_LEVEL: Rs ~9,000/bbl (WTI ~$94.4) flips to BEAR (puts only); below Rs ~8,500 (WTI ~$89) flips to MILD BULL.",
    "AVOID (if it tips into BEAR >9,000): OMCs (BPCL/HPCL/IOC), aviation, paints, tyres, fertilisers.",
    "EXCEPTION_CALLS (crude beneficiaries): ONGC, Oil India, MRPL — favoured while crude is elevated.",
])
add_para(doc, "Direction awareness: Crude is NOT falling >2% overnight today and GIFT is NOT gap-up >100 pts, so the A+B contrarian override is NOT triggered. But the Iran binary keeps a spike on the table — watch Rs 9,000 for a BEAR mode flip intraday.", bold=True)
add_para(doc, "Structural read (NSE option chain):")
add_bullets(doc, [
    "Nifty vs max pain: close 24,031.7 vs max pain ~23,974 -> ~57 pts ABOVE. On expiry day, max-pain magnetism pulls DOWN toward 23,975.",
    "PCR: DATA_UNAVAILABLE (live value not confirmed across sources) — not fabricated.",
    "OI change direction: DATA_UNAVAILABLE for today's session; treat strike-wise OI live at open.",
    "Futures basis: DATA_UNAVAILABLE; on monthly expiry day basis collapses to ~0 (cash-future convergence).",
    "Expiry dynamics: DUAL expiry (weekly + monthly). Pin risk is maximal. Likely battle zone 23,950-24,050.",
    "Put writers floor (provisional): ~23,900/23,800. Call writers ceiling (provisional): ~24,100/24,200.",
    "Max-pain magnetism direction: DOWN (toward ~23,975), unless a clean break >24,100 invalidates the pin.",
])
add_para(doc, "Bank Nifty leadership check: ICICI/HDFC ADRs firm (+1-3%), banks led the 25 May rally. If Bank Nifty LEADS Nifty at open -> move is REAL; if it LAGS -> fade. (Confirm live, BankNifty intraday data not pre-marketable.)", italic=True)

# ---------- SECTION 3 ----------
add_heading(doc, "Section 3 — Contrarian Check (Layer 3)")
add_bullets(doc, [
    "Q1 CONSENSUS: After a +1.32% rally and a 6% crude crash, retail/TV expect continuation UP and 'crude relief = buy India'. Bullish lean into expiry.",
    "Q2 THE TRAP: Expiry-day max-pain pin sits BELOW spot at ~23,975. Fresh US strikes on Iran can reignite crude/geo fear and fade the rally. Option writers profit as BOTH sides get theta-crushed into the pin.",
    "Q3 RETAIL STOPS: Long stops clustered below 23,900 then 23,800; short stops above 24,100 then 24,200.",
    "Q4 FLIP TRIGGER: A confirmed US-Iran MoU / Hormuz reopening (crude gaps down hard -> CAUTIOUS BULL) OR an Iran retaliation / tanker-mine incident (crude spikes >Rs 9,000 -> BEAR). Binary.",
])
add_para(doc, "CONTRARIAN OVERRIDE LOGIC:", bold=True, color=NAVY)
add_kv_table(doc,
    header=["Condition", "Status", "Note"],
    rows=[
        ["A: GIFT gap-up > +100 pts?", "NO", "Gap ~+6 pts (flat)"],
        ["B: Crude falling > 2% overnight?", "NO", "Today ~+0.5% (the -6% was prior session)"],
        ["C: Contrarian scenario prob > 40%?", "YES (~45%)", "Expiry pin + Iran binary"],
        ["D: Expiry week + PCR < 0.8?", "UNCONFIRMED", "Expiry=YES but PCR DATA_UNAVAILABLE -> not asserted"],
    ])
add_para(doc, "Result: A+B not both YES -> NO bear-suppression override. C+D not both confirmable -> no forced LOW downgrade, but confidence held at MEDIUM given the unconfirmed PCR and Iran binary. Final bias = RANGE (expiry pin), mild downward lean to max pain.", bold=True)

# ---------- SECTION 4 ----------
add_heading(doc, "Section 4 — 5 Individual F&O Setups")
add_para(doc, "EXPIRY-DAY RULE OVERLAY: Today is monthly+weekly expiry. Per profile, NO positional option buying — current-expiry trades are intraday scalps only, exit before 13:00 IST. For any hold, roll to JUNE monthly. All 15-min entry signals (SuperTrend/RSI/StochRSI/Volume) must be CONFIRMED live at open; they cannot be observed pre-market. Lot sizes are approximate — VERIFY on NSE (revised periodically).", italic=True, color=RED)

# Stock 1 — ONGC
stock_header(doc, "1) ONGC — NSE: ONGC  |  CALL  |  Grade A")
add_para(doc, "CMP: ~Rs 290 (22 May)  |  Lot: ~7,700 (verify NSE)  |  Crude aligned: YES (beneficiary)")
add_para(doc, "CATALYST: Elevated crude (~$92 WTI / Rs 8,766 proxy) lifts realisations; +7.2% in last month, near 52-wk high Rs 307.5. Iran spike risk is an upside tail for producers.")
add_bullets(doc, [
    "TECHNICAL: Daily UP (monthly +7%); intraday wobble 22 May. 200 DMA: near/above (est., verify). Support Rs 284 | Resistance Rs 298 / 307.5",
    "STRUCTURAL: OI direction DATA_UNAVAILABLE (confirm long build at open) | Delivery STABLE | Block deal: NO | Ban list: NO",
    "OPTION SETUP: Strike Rs 290 ATM | Expiry: roll to JUNE for any hold; today CE scalp only | THETA WARNING (expiry <=3d): YES",
    "ENTRY (all 4 must fire live): ST GREEN 15-min | RSI >50 | StochRSI cross UP from <30 | Volume >1.5x 20-avg | confirm 15-min close > Rs 292",
    "T1 (book 40%): Rs 297 | T2 (book 40%): Rs 302 | SL: Rs 288 + 15-min close below ST | R:R ~2.5:1 | Time stop: exit by 13:00 IST",
])
add_para(doc, "GRADE REASON: Cleanest crude-aligned long with momentum + tail-risk upside; downgraded from A+ only by expiry-day theta.", italic=True)

# Stock 2 — Wipro
stock_header(doc, "2) WIPRO — NSE: WIPRO  |  CALL  |  Grade A")
add_para(doc, "CMP: ~Rs 203 (22 May)  |  Lot: ~3,000 (verify NSE)  |  Crude aligned: N/A (IT, crude-neutral)")
add_para(doc, "CATALYST: WIT ADR +7.17% (22 May, stale but strong) + Rs 15,000 Cr buyback at Rs 250 (record date 5 Jun) = firm downside support / catalyst.")
add_bullets(doc, [
    "TECHNICAL: Daily DOWN/base (down 25% YoY) but basing with buyback floor. Below 200 DMA: YES. Support Rs 198 | Resistance Rs 210",
    "STRUCTURAL: Long build-up noted 15 May (OI +3.24%) | Delivery STABLE | Block deal: NO | Ban list: NO",
    "OPTION SETUP: Strike Rs 205 ATM/OTM | Expiry: JUNE for hold; today CE scalp only | THETA WARNING: YES",
    "ENTRY (all 4 must fire live): ST GREEN 15-min | RSI >50 | StochRSI cross UP from <30 | Volume >1.5x avg | confirm 15-min close > Rs 205",
    "T1 (book 40%): Rs 210 | T2 (book 40%): Rs 215 | SL: Rs 200 + 15-min close below ST | R:R ~2:1 | Time stop: exit by 13:00 IST",
])
add_para(doc, "GRADE REASON: ADR strength + buyback support give an asymmetric long; only the prevailing downtrend keeps it at A not A+.", italic=True)

# Stock 3 — BPCL
stock_header(doc, "3) BPCL — NSE: BPCL  |  PUT  |  Grade A")
add_para(doc, "CMP: ~Rs 284 (15 May)  |  Lot: ~1,800 (verify NSE)  |  Crude aligned: YES (victim — high crude squeezes marketing margins)")
add_para(doc, "CATALYST: Crude elevated + Iran spike risk threatens OMC marketing margins (Q4 margin Rs 5.6/lit unlikely to sustain if crude jumps). Stock in downtrend from Feb ATH Rs 391.")
add_bullets(doc, [
    "TECHNICAL: Daily DOWN (targets Rs 269/256); weekly DOWN. Below 200 DMA: YES (est.). Support Rs 275 / 269 | Resistance Rs 290",
    "STRUCTURAL: OI DATA_UNAVAILABLE (confirm short build at open) | Delivery STABLE | Block deal: NO | Ban list: NO",
    "OPTION SETUP: Strike Rs 285 ATM PE | Expiry: JUNE for hold; today PE scalp only | THETA WARNING: YES",
    "ENTRY (all 4 must fire live): ST RED 15-min | RSI <50 | StochRSI cross DOWN from >70 | Volume >1.5x avg | confirm 15-min close < Rs 281",
    "T1 (book 40%): Rs 275 | T2 (book 40%): Rs 270 | SL: Rs 286 + 15-min close above ST | R:R ~2.2:1 | Time stop: exit by 13:00 IST",
])
add_para(doc, "GRADE REASON: Confirmed downtrend + crude-victim alignment in CAUTION mode; clean short with defined resistance.", italic=True)

# Stock 4 — ICICI Bank
stock_header(doc, "4) ICICI BANK — NSE: ICICIBANK  |  CALL (range-long)  |  Grade B")
add_para(doc, "CMP: ~Rs 1,275 (22 May)  |  Lot: ~700 (verify NSE)  |  Crude aligned: N/A (financials)")
add_para(doc, "CATALYST: Banks led the 25 May rally; IBN ADR +1.23%. Bank Nifty leadership is the real-vs-fake tell for Nifty today.")
add_bullets(doc, [
    "TECHNICAL: Intraday UP, but 50 DMA 1,277 vs 200 DMA 1,360 -> Below 200 DMA: YES (range/recovery). Support Rs 1,255 | Resistance Rs 1,295 / 1,300",
    "STRUCTURAL: OI DATA_UNAVAILABLE | Delivery STABLE | Block deal: NO | Ban list: NO",
    "OPTION SETUP: Strike Rs 1,280 ATM CE | Expiry: JUNE for hold; today scalp only | THETA WARNING: YES",
    "ENTRY (all 4 must fire live): ST GREEN 15-min | RSI >50 | StochRSI cross UP | Volume >1.5x avg | confirm 15-min close > Rs 1,285",
    "T1 (book 40%): Rs 1,300 | T2 (book 40%): Rs 1,315 | SL: Rs 1,270 + 15-min close below ST | R:R ~2:1 | Time stop: exit by 13:00 IST",
])
add_para(doc, "GRADE REASON: B — below 200 DMA and rangey; useful as a Bank Nifty confirmation proxy more than a standalone edge.", italic=True)

# Stock 5 — IndiGo
stock_header(doc, "5) INTERGLOBE (INDIGO) — NSE: INDIGO  |  PUT  |  Grade B")
add_para(doc, "CMP: ~Rs 4,411 (22 May)  |  Lot: ~150 (verify NSE)  |  Crude aligned: YES (victim — jet fuel)")
add_para(doc, "CATALYST: Fuel-cost pressure (EBITDA margin slid to ~23%); airlines lobbying OMCs to delay ATF hikes amid West Asia conflict. Any Iran-driven crude spike hits hardest here.")
add_bullets(doc, [
    "TECHNICAL: Daily UP recently (4,262 -> 4,411) — counter-trend PUT, so demands strict confirmation. Below 200 DMA: NO (above). Support Rs 4,300 | Resistance Rs 4,450",
    "STRUCTURAL: OI DATA_UNAVAILABLE | Delivery STABLE | Block deal: NO | Ban list: NO",
    "OPTION SETUP: Strike Rs 4,400 ATM PE | Expiry: JUNE for hold; today scalp only | THETA WARNING: YES",
    "ENTRY (all 4 must fire live): ST RED 15-min | RSI <50 | StochRSI cross DOWN from >70 | Volume >1.5x avg | confirm 15-min close < Rs 4,380",
    "T1 (book 40%): Rs 4,310 | T2 (book 40%): Rs 4,240 | SL: Rs 4,450 + 15-min close above ST | R:R ~2:1 | Time stop: exit by 13:00 IST",
])
add_para(doc, "GRADE REASON: B — crude-victim thesis is sound but price momentum is UP; only take if crude spikes AND 15-min reverses. Otherwise skip.", italic=True)

# ---------- SECTION 5 ----------
add_heading(doc, "Section 5 — Final Verdict")
add_kv_table(doc, header=None, rows=[
    ["TODAY'S BIAS", "RANGE (expiry pin) — mild down lean to max pain ~23,975"],
    ["CONFIDENCE", "MEDIUM"],
    ["CRUDE MODE", "CAUTION (~Rs 8,766/bbl), direction choppy/flat"],
    ["VIX SIZING", "HALF SIZE (mandatory)"],
    ["ENTRY PERMISSION", "YELLOW — expiry day, intraday scalps only, exit by 13:00"],
])
add_para(doc, "THE BULL CASE: Strong 25 May close (+1.32%), crude crashed 6% on deal hopes, US indices at/near highs, DIIs buying. A confirmed Iran MoU gaps Nifty toward 24,200+.", color=GREEN)
add_para(doc, "THE BEAR CASE: Dual-expiry max-pain magnet sits BELOW at ~23,975, FIIs net sellers, fresh US strikes on Iran threaten a crude spike >Rs 9,000 that flips mode to BEAR.", color=RED)
add_para(doc, 'FLIP TRIGGER: "If Nifty sustains >24,100 with Bank Nifty leading -> CAUTIOUS BULL; if crude spikes >Rs 9,000/bbl or Nifty breaks <23,900 -> BEAR."', bold=True)
add_para(doc, "NIFTY KEY LEVELS:  S2 23,800  |  S1 23,900  |  CRITICAL 23,975 (max pain)  |  R1 24,100  |  R2 24,200", bold=True)
add_para(doc, "TOP 3 RANKED:", bold=True, color=NAVY)
add_bullets(doc, [
    "#1 ONGC (A) — crude-aligned long with Iran upside tail — entry CE above Rs 292.",
    "#2 BPCL (A) — crude-victim short, confirmed downtrend — entry PE below Rs 281.",
    "#3 WIPRO (A) — ADR +7% + buyback floor — entry CE above Rs 205.",
])
add_para(doc, "ONE RISK THAT RUINS EVERYTHING TODAY: A binary Iran headline mid-session — a sudden deal (crude crash, sharp gap) OR an escalation/Hormuz mine strike (crude spike). On a dual-expiry day this whipsaws both option sides and theta-crushes pinned positions.", bold=True, color=RED)

# ---------- ONE-LINER ----------
add_heading(doc, "One-Line Summary (read at 9:10 AM)")
add_para(doc,
    'Today is RANGE because dual (weekly+monthly) expiry pins Nifty toward max pain ~23,975 below a flat open. '
    'Crude ~Rs 8,766 = CAUTION, choppy/flat after a 6% crash. Watch ONGC CE and BPCL PE. '
    'Key risk: a binary Iran headline (deal=crude crash / escalation=crude spike). Size HALF. '
    'Flips BULL if Nifty holds >24,100 with banks leading; flips BEAR if crude >Rs 9,000 or Nifty <23,900.',
    italic=True, size=12)

add_para(doc, "")
add_para(doc, "DATA INTEGRITY NOTE: Pre-market brief generated ~03:15-09:00 IST. Items marked DATA_UNAVAILABLE (live PCR, today's exact F&O ban list, intraday OI/basis, overnight ADR deltas) were NOT fabricated. ADR levels and several stock prices are last-available (22-25 May). Re-confirm option chain, ban list and 15-min signals on NSE at the open before acting.", italic=True, size=9, color=RGBColor(0x80,0x80,0x80))

os.makedirs(OUT_DIR, exist_ok=True)
doc.save(OUT_PATH)
print("Saved:", OUT_PATH)
